from docx import Document
import os

files = os.listdir(r'D:\code\ideaProject\petService\output\word_docs')
print("Files:", files)

for fname in files:
    if fname.endswith('.docx'):
        fpath = os.path.join(r'D:\code\ideaProject\petService\output\word_docs', fname)
        print(f"\n=== {fname} ===")
        try:
            doc = Document(fpath)
            print(f"Paragraphs: {len(doc.paragraphs)}")
            for i, p in enumerate(doc.paragraphs[:15]):
                text = p.text.strip()
                if text:
                    print(f"  {i}: [{p.style.name}] {text[:150]}")
        except Exception as e:
            print(f"Error: {e}")